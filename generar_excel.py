import csv
import openpyxl
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter

# Crear libro nuevo y limpio
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "evidencias"

# Asegurar cuadrícula visible
ws.views.sheetView[0].showGridLines = True

# Borde estándar negro/gris para todas las celdas
thin_side = Side(style='thin', color='000000')
all_borders = Border(
    left=thin_side,
    right=thin_side,
    top=thin_side,
    bottom=thin_side
)

# Leer datos directamente desde evidencias.csv
csv_path = r"c:\Users\Leonardo\Desktop\cloud-forge-ci\evidencias.csv"
with open(csv_path, mode="r", encoding="utf-8") as f:
    reader = csv.reader(f)
    for row in reader:
        if row: # evitar filas vacías si las hay
            ws.append(row)

# Aplicar únicamente 'todos los bordes' a cada celda con datos, sin colores ni estilos extra
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
    for cell in row:
        cell.border = all_borders

# Autoajustar ancho de columnas para que se lea el texto con comodidad
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        val = str(cell.value or '')
        if len(val) > max_len:
            max_len = len(val)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

output_path = r"c:\Users\Leonardo\Desktop\cloud-forge-ci\Laboratorio_2_Evidencias.xlsx"
wb.save(output_path)
print("Excel generado sin formato y con todos los bordes en:", output_path)
